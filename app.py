import streamlit as st
import pandas as pd
import re
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── helpers ───────────────────────────────────────────────────────────────────

def to_title_case(text: str) -> str:
    s = str(text).strip()
    return ' '.join(word.capitalize() for word in s.split())


def convert_column_type(series: pd.Series, dtype: str) -> pd.Series:
    if dtype == "Integer":
        def to_int(v):
            try: return int(float(str(v).strip()))
            except: return v
        return series.apply(lambda v: to_int(v) if pd.notna(v) else v)
    elif dtype == "Float":
        def to_float(v):
            try: return float(str(v).strip())
            except: return v
        return series.apply(lambda v: to_float(v) if pd.notna(v) else v)
    elif dtype == "Phone (last 10 digits)":
        def to_phone(v):
            digits = re.sub(r"\D", "", str(v))
            return digits[-10:] if len(digits) >= 10 else digits
        return series.apply(lambda v: to_phone(v) if pd.notna(v) else v)
    return series


def read_file(uploaded) -> pd.DataFrame:
    name = uploaded.name.lower()
    if name.endswith('.csv'):
        raw = uploaded.read()
        uploaded.seek(0)
        sample = raw[:4096].decode('utf-8', errors='replace')
        sep = ','
        for candidate in [',', ';', '\t', '|']:
            if sample.count(candidate) > sample.count(sep):
                sep = candidate
        text = raw.decode('utf-8', errors='replace')
        lines = text.splitlines()
        header = lines[0].split(sep) if lines else []
        max_cols = max((len(l.split(sep)) for l in lines if l.strip()), default=len(header))
        extra = max_cols - len(header)
        col_names = header + [f'_extra_{i}' for i in range(extra)]
        df = pd.read_csv(
            io.BytesIO(raw), sep=sep, on_bad_lines='skip',
            engine='python', index_col=False, names=col_names, skiprows=1,
        )
        df = df[[c for c in df.columns if not c.startswith('_extra_')]]
    else:
        df = pd.read_excel(uploaded, index_col=None)

    unnamed_mask = df.columns.astype(str).str.match(r'^Unnamed')
    all_empty_mask = df.isnull().all(axis=0)
    df = df.loc[:, ~(unnamed_mask & all_empty_mask)]
    df = df.dropna(how='all')
    df = df.reset_index(drop=True)
    return df


def df_to_xlsx_bytes(df: pd.DataFrame, col_padding: int = 0) -> bytes:
    wb_out = Workbook()
    ws = wb_out.active
    header_font = Font(bold=True, name="Calibri", size=12, color="000000")
    data_font   = Font(bold=False, name="Calibri", size=11, color="000000")
    header_fill  = PatternFill("solid", fgColor="FFFF00")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_align   = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    n_rows = len(df)
    n_cols = len(df.columns)
    for ci, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = header_align; cell.border = cell_border
    for ri in range(n_rows):
        for ci in range(n_cols):
            val = df.iloc[ri, ci]
            if hasattr(val, 'item'): val = val.item()
            cell = ws.cell(row=ri + 2, column=ci + 1, value=val)
            cell.font = data_font; cell.fill = white_fill
            cell.border = cell_border; cell.alignment = data_align
    for ci, col in enumerate(df.columns, start=1):
        col_letter = get_column_letter(ci)
        try:
            data_max = df[col].dropna().astype(str).str.len().max()
            data_max = int(data_max) if pd.notna(data_max) else 0
        except: data_max = 0
        max_len = max(len(str(col)), data_max)
        ws.column_dimensions[col_letter].width = max_len + col_padding + 2
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return buf.read()


def df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode('utf-8')


def _make_error_pdf(message: str) -> bytes:
    """Build a bare-minimum valid PDF containing a single text message."""
    msg = message.encode('latin-1', errors='replace')
    stream = (
        b"BT /F1 12 Tf 50 750 Td (" + msg + b") Tj ET"
    )
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842]"
         b" /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"),
        b"4 0 obj\n<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream + b"\nendstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for obj in objects:
        offsets.append(len(body))
        body += obj
    xref_pos = len(body)
    xref = b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    xref += b"0000000000 65535 f \n"
    for off in offsets:
        xref += str(off).zfill(10).encode() + b" 00000 n \n"
    trailer = (b"trailer\n<< /Size " + str(len(objects) + 1).encode()
               + b" /Root 1 0 R >>\nstartxref\n"
               + str(xref_pos).encode() + b"\n%%EOF\n")
    return body + xref + trailer


def df_to_pdf_bytes(df: pd.DataFrame) -> bytes:
    """
    Export dataframe to PDF using reportlab with Paragraph cells so text
    wraps instead of overflowing.  Falls back to a valid error PDF when
    reportlab is not installed, so the browser never receives raw bytes.
    """
    try:
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER
    except ImportError:
        return _make_error_pdf(
            "PDF export requires the reportlab package. "
            "Install it with:  pip install reportlab"
        )

    page_w, page_h = landscape(A4)
    margin = 20
    usable_w = page_w - 2 * margin

    col_count = len(df.columns)
    col_width = usable_w / col_count if col_count else usable_w

    hdr_style = ParagraphStyle(
        "hdr", fontName="Helvetica-Bold", fontSize=8,
        alignment=TA_CENTER, leading=10, textColor=colors.black,
        wordWrap="CJK",
    )
    cell_style = ParagraphStyle(
        "cell", fontName="Helvetica", fontSize=7,
        alignment=TA_CENTER, leading=9, textColor=colors.black,
        wordWrap="CJK",
    )

    def safe_str(v):
        try:
            if pd.isna(v):
                return ""
        except (TypeError, ValueError):
            pass
        return str(v)

    header_row = [Paragraph(safe_str(c), hdr_style) for c in df.columns]
    data_rows  = [
        [Paragraph(safe_str(df.iloc[ri, ci]), cell_style) for ci in range(col_count)]
        for ri in range(len(df))
    ]
    table_data = [header_row] + data_rows

    table = Table(
        table_data,
        colWidths=[col_width] * col_count,
        repeatRows=1,
        hAlign="CENTER",
    )
    table.setStyle(TableStyle([
        ('BACKGROUND',    (0, 0), (-1, 0),  colors.HexColor('#FFFF00')),
        ('TEXTCOLOR',     (0, 0), (-1, 0),  colors.black),
        ('ALIGN',         (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND',    (0, 1), (-1, -1), colors.white),
        ('GRID',          (0, 0), (-1, -1), 0.4, colors.black),
        ('TOPPADDING',    (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING',   (0, 0), (-1, -1), 4),
        ('RIGHTPADDING',  (0, 0), (-1, -1), 4),
    ]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin,
    )
    doc.build([table])
    buf.seek(0)
    return buf.read()


def df_to_html_preview(df: pd.DataFrame) -> str:
    """Render the full dataframe as a scrollable HTML table."""
    rows_html = ""
    for ri in range(len(df)):
        cells = ""
        for ci in range(len(df.columns)):
            val = df.iloc[ri, ci]
            val = "" if pd.isna(val) else str(val)
            cells += f'<td>{val}</td>'
        rows_html += f"<tr>{cells}</tr>"
    headers = "".join(f"<th>{col}</th>" for col in df.columns)
    return f"""
    <div style="overflow-x:auto; margin-top:8px; max-height:600px; overflow-y:auto;">
    <table style="border-collapse:collapse; font-family:Calibri,sans-serif; width:100%; font-size:13px;">
        <thead style="position:sticky; top:0; z-index:1;">
            <tr>{headers}</tr>
        </thead>
        <tbody>{rows_html}</tbody>
    </table>
    </div>
    <style>
    table th {{
        background:#FFFF00 !important; color:#000 !important;
        font-weight:bold !important; font-size:14px !important;
        text-align:center !important; padding:7px 10px !important;
        border:1.5px solid #000 !important; white-space:nowrap;
        position:sticky; top:0;
    }}
    table td {{
        background:#FFFFFF !important; color:#000 !important;
        font-size:13px !important; text-align:center !important;
        padding:5px 10px !important; border:1px solid #000 !important;
    }}
    </style>
    """


def render_sort_ui(available_cols: list, key_prefix: str):
    """
    Renders a dynamic sort UI where the user adds sort levels one by one.
    Returns (sort_cols_ordered, sort_orders_dict).
    sort_cols_ordered — list of column names in priority order (1st = primary key)
    sort_orders_dict  — {col: True(asc) / False(desc)}
    """
    st.markdown('<div class="tip-box">Add sort levels in priority order. The <b>first</b> row is the primary sort key, the second is the tie-breaker, and so on. Each column defaults to <b>Ascending ↑</b>.</div>', unsafe_allow_html=True)

    num_sorts = st.number_input(
        "Number of sort levels",
        min_value=0, max_value=min(10, len(available_cols)),
        value=0, step=1, key=f"{key_prefix}_num_sorts"
    )

    sort_cols   = []
    sort_orders = {}

    if int(num_sorts) > 0:
        # Table-style header row
        hc1, hc2, hc3 = st.columns([0.5, 2, 1.5])
        hc1.markdown("**Level**")
        hc2.markdown("**Column**")
        hc3.markdown("**Order**")

        already_chosen = []
        for i in range(int(num_sorts)):
            rc1, rc2, rc3 = st.columns([0.5, 2, 1.5])
            rc1.markdown(f"<div style='padding-top:8px; font-weight:600; color:#475569;'>#{i+1}</div>", unsafe_allow_html=True)

            # Exclude columns already picked at a higher priority level
            remaining = [c for c in available_cols if c not in already_chosen]
            options = ["— select —"] + remaining

            sort_col = rc2.selectbox(
                f"Sort level {i+1}",
                options=options,
                index=0,
                key=f"{key_prefix}_sort_col_{i}",
                label_visibility="collapsed"
            )
            sort_dir = rc3.selectbox(
                f"Direction {i+1}",
                options=["Ascending ↑", "Descending ↓"],
                index=0,
                key=f"{key_prefix}_sort_dir_{i}",
                label_visibility="collapsed"
            )

            if sort_col != "— select —":
                sort_cols.append(sort_col)
                sort_orders[sort_col] = (sort_dir == "Ascending ↑")
                already_chosen.append(sort_col)

    return sort_cols, sort_orders


def apply_transformations(df, drop_empty, type_conversions, data_camel_cols, do_title_headers,
                           sort_cols, sort_orders):
    result = df.copy()
    if drop_empty:
        result = result.dropna(how='all')
    for col, dtype in type_conversions.items():
        if col in result.columns:
            result[col] = convert_column_type(result[col], dtype)
    for col in data_camel_cols:
        if col in result.columns:
            result[col] = result[col].apply(
                lambda v: to_title_case(str(v)) if pd.notna(v) else v)
    # Sorting — apply before rename so col names still match
    if sort_cols:
        ascending_flags = [sort_orders.get(c, True) for c in sort_cols]
        try:
            result = result.sort_values(by=sort_cols, ascending=ascending_flags).reset_index(drop=True)
        except Exception:
            pass
    if do_title_headers:
        result.columns = [to_title_case(c) for c in result.columns]
    return result


def download_section(result, col_padding, base_name, key_prefix):
    fmt = st.selectbox(
        "Download format",
        options=["xlsx (default)", "csv", "pdf", "xls", "xlsm"],
        index=0,
        key=f"{key_prefix}_fmt",
    )

    if fmt == "xlsx (default)":
        data   = df_to_xlsx_bytes(result, col_padding)
        fname  = f"{base_name}.xlsx"
        mime   = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        label  = "⬇️  Download .xlsx"
    elif fmt == "csv":
        data  = df_to_csv_bytes(result)
        fname = f"{base_name}.csv"
        mime  = "text/csv"
        label = "⬇️  Download .csv"
    elif fmt == "pdf":
        data  = df_to_pdf_bytes(result)
        fname = f"{base_name}.pdf"
        mime  = "application/pdf"
        label = "⬇️  Download .pdf"
    elif fmt == "xls":
        data  = df_to_xlsx_bytes(result, col_padding)
        fname = f"{base_name}.xls"
        mime  = "application/vnd.ms-excel"
        label = "⬇️  Download .xls"
    elif fmt == "xlsm":
        data  = df_to_xlsx_bytes(result, col_padding)
        fname = f"{base_name}.xlsm"
        mime  = "application/vnd.ms-excel.sheet.macroEnabled.12"
        label = "⬇️  Download .xlsm"

    btn_class = "map-download-btn" if "map" in key_prefix else ""
    st.markdown(f'<div class="{btn_class}">', unsafe_allow_html=True)
    st.download_button(label=label, data=data, file_name=fname, mime=mime, key=f"{key_prefix}_dl")
    st.markdown('</div>', unsafe_allow_html=True)


# ── page config ───────────────────────────────────────────────────────────────

st.set_page_config(page_title="Sheet Toolkit", page_icon="📊", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600;700&family=JetBrains+Mono:wght@400;500&display=swap');
html, body, [class*="css"] { font-family: 'Space Grotesk', sans-serif; }
.block-container { padding-top: 1.5rem; }
h1 { font-size: 2rem !important; font-weight: 700 !important; letter-spacing: -0.5px; }
h2 { font-size: 1.35rem !important; font-weight: 600 !important; }
h3 { font-size: 1.1rem !important; font-weight: 600 !important; }

/* tab styling */
.stTabs [data-baseweb="tab-list"] {
    gap: 6px; background: #e2e8f0; padding: 5px; border-radius: 12px;
}
.stTabs [data-baseweb="tab"] {
    border-radius: 8px; padding: 10px 28px;
    font-weight: 700; font-size: 0.95rem; color: #1e293b !important;
}
.stTabs [data-baseweb="tab"]:hover { background: #cbd5e1 !important; color: #0f172a !important; }
.stTabs [aria-selected="true"] { background: #ffffff !important; color: #0f172a !important; box-shadow: 0 1px 4px rgba(0,0,0,0.15); }
.stTabs [data-baseweb="tab"] p,
.stTabs [data-baseweb="tab"] span,
.stTabs [data-baseweb="tab"] div { color: #1e293b !important; font-weight: 700 !important; }

/* download button — default */
.stDownloadButton > button {
    background: #2E4057 !important; color: white !important;
    border-radius: 8px !important; font-weight: 600 !important;
    padding: 0.5rem 1.5rem !important; border: none !important;
}
.stDownloadButton > button:hover { background: #1a2d42 !important; }

/* mapping download button */
.map-download-btn .stDownloadButton > button {
    background: #16A34A !important; color: white !important;
    border-radius: 10px !important; font-weight: 700 !important;
    font-size: 1rem !important; padding: 0.75rem 2.5rem !important;
    border: none !important; box-shadow: 0 4px 14px rgba(22,163,74,0.4) !important;
}
.map-download-btn .stDownloadButton > button:hover {
    background: #15803D !important; box-shadow: 0 6px 18px rgba(22,163,74,0.55) !important;
}

/* info boxes */
.tip-box {
    background: #EFF6FF; border-left: 4px solid #3B82F6;
    padding: 0.75rem 1rem; border-radius: 0 8px 8px 0;
    font-size: 0.88rem; color: #1E40AF; margin-bottom: 1rem;
}
code { font-family: 'JetBrains Mono', monospace; font-size: 0.82rem; }

/* sort level table header */
.sort-header { font-weight: 700; color: #475569; font-size: 0.85rem; padding-bottom: 4px; }
</style>
""", unsafe_allow_html=True)

# ── sidebar ───────────────────────────────────────────────────────────────────

with st.sidebar:
    st.markdown("## 📊 Sheet Toolkit")
    st.markdown("---")
    st.markdown("""
**Two modes:**

🔧 **Structure** — Clean & format a single file  
🔗 **Mapping** — Map a base file to an expected schema
    """)
    st.markdown("---")
    st.markdown("**Supported formats:** `.csv`, `.xlsx`, `.xls`, `.xlsm`")
    st.markdown("---")
    st.caption("Tip: Downloads available in xlsx, csv, pdf, xls, xlsm.")

# ── main ──────────────────────────────────────────────────────────────────────

st.markdown("# 📊 Sheet Toolkit")
st.markdown("Clean, structure, and map your spreadsheet data in seconds.")

tab1, tab2 = st.tabs(["🔧  Structure File", "🔗  Mapping"])


# ═══════════════════════════════════════════════════════════════════════════════
# TAB 1 — STRUCTURE FILE
# ═══════════════════════════════════════════════════════════════════════════════

with tab1:
    st.markdown("### Upload your file")
    uploaded = st.file_uploader("Drag & drop or browse", type=["csv","xlsx","xls","xlsm"], key="struct_upload")

    if uploaded:
        df = read_file(uploaded)
        original_cols = list(df.columns)

        st.success(f"Loaded **{len(df)} rows × {len(df.columns)} cols**")

        with st.expander("👀 Preview raw data", expanded=False):
            st.dataframe(df.head(10), use_container_width=True, hide_index=True)

        st.markdown("---")
        st.markdown("### ⚙️ Transformations")

        col_a, col_b = st.columns(2)
        with col_a:
            st.markdown("#### 1️⃣ Title Case headers")
            st.markdown('<div class="tip-box">All column headers will be converted to Title Case (e.g. "FULL NAME" → "Full Name").</div>', unsafe_allow_html=True)
            do_all_camel = st.checkbox("Convert ALL headers to Title Case", value=True, key="all_camel")

            st.markdown("#### 2️⃣ Title Case selected data columns")
            st.markdown('<div class="tip-box">Pick columns whose cell values should be converted to Title Case (e.g. RAJ NANGARE → Raj Nangare).</div>', unsafe_allow_html=True)
            data_camel_cols = st.multiselect("Select columns", options=original_cols, key="data_camel")

        with col_b:
            st.markdown("#### 3️⃣ Column width padding")
            st.markdown('<div class="tip-box">Extra characters added to each column beyond its longest value.</div>', unsafe_allow_html=True)
            col_padding = st.slider("Extra padding (chars)", 0, 20, 5, key="col_pad")

            st.markdown("#### 4️⃣ Drop empty rows")
            drop_empty = st.checkbox("Drop fully-empty rows", value=False, key="drop_empty")

        st.markdown("---")
        st.markdown("#### 5️⃣ Convert column data types")
        st.markdown('<div class="tip-box">Select columns to convert. Phone number keeps only the last 10 digits.</div>', unsafe_allow_html=True)
        type_conversions = {}
        num_conversions = st.number_input("How many columns to convert?", min_value=0, max_value=20, value=0, step=1, key="num_conv")
        for i in range(int(num_conversions)):
            cc1, cc2 = st.columns([2, 1])
            col_choice = cc1.selectbox(f"Column #{i+1}", options=["— select —"] + original_cols, key=f"conv_col_{i}")
            type_choice = cc2.selectbox("Convert to", options=["Integer", "Float", "Phone (last 10 digits)"], key=f"conv_type_{i}")
            if col_choice != "— select —":
                type_conversions[col_choice] = type_choice

        st.markdown("---")
        st.markdown("#### 6️⃣ Sort data")
        sort_cols, sort_orders = render_sort_ui(original_cols, key_prefix="struct")

        st.markdown("---")
        st.markdown("### 🔄 Preview result")

        result = apply_transformations(
            df, drop_empty, type_conversions, data_camel_cols,
            do_all_camel, sort_cols, sort_orders
        )

        st.markdown(f"Showing all **{len(result)}** rows")
        st.markdown(df_to_html_preview(result), unsafe_allow_html=True)

        st.markdown("---")
        base_name = uploaded.name.rsplit('.', 1)[0] + "_structured"
        download_section(result, col_padding, base_name, key_prefix="struct")


# ═══════════════════════════════════════════════════════════════════════════════
# TAB 2 — MAPPING
# ═══════════════════════════════════════════════════════════════════════════════

with tab2:
    st.markdown("### Upload both files")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**📋 Expected / Target schema file**")
        st.caption("Defines the column names you want in the output.")
        exp_file = st.file_uploader("Expected file", type=["csv","xlsx","xls","xlsm"], key="exp_upload", label_visibility="collapsed")
    with c2:
        st.markdown("**📂 Base / Source data file**")
        st.caption("The file whose data will be mapped to the expected schema.")
        base_file = st.file_uploader("Base file", type=["csv","xlsx","xls","xlsm"], key="base_upload", label_visibility="collapsed")

    if exp_file and base_file:
        exp_df  = read_file(exp_file)
        base_df = read_file(base_file)
        exp_cols  = list(exp_df.columns)
        base_cols = list(base_df.columns)

        st.success(f"Expected: **{len(exp_cols)} cols** | Base: **{len(base_df)} rows × {len(base_cols)} cols**")

        with st.expander("👀 Preview expected schema (first 5 rows)", expanded=False):
            st.dataframe(exp_df.head(5), use_container_width=True, hide_index=True)
        with st.expander("👀 Preview base data (first 5 rows)", expanded=False):
            st.dataframe(base_df.head(5), use_container_width=True, hide_index=True)

        st.markdown("---")
        st.markdown("### 🗺️ Column Mapping")
        st.markdown('<div class="tip-box">For each expected column, choose the matching column from your base file — or leave as <code>— (empty column) —</code> to include the column header with no data.</div>', unsafe_allow_html=True)

        mapping = {}
        EMPTY_OPT = "— (empty column) —"
        base_options = [EMPTY_OPT] + base_cols

        def auto_match(exp_col):
            for bc in base_cols:
                if bc.strip().lower() == exp_col.strip().lower():
                    return bc
            return EMPTY_OPT

        cols_per_row = 3
        for i in range(0, len(exp_cols), cols_per_row):
            row_cols = exp_cols[i:i + cols_per_row]
            grid = st.columns(len(row_cols))
            for j, exp_col in enumerate(row_cols):
                default = auto_match(exp_col)
                chosen = grid[j].selectbox(f"`{exp_col}`", options=base_options,
                                           index=base_options.index(default), key=f"map_{i}_{j}")
                mapping[exp_col] = None if chosen == EMPTY_OPT else chosen

        st.markdown("---")
        st.markdown("### ⚙️ Output options")
        mc1, mc2 = st.columns(2)
        with mc1:
            map_all_camel = st.checkbox("Convert all output headers to Title Case", value=True, key="map_camel")
            map_data_camel_cols = st.multiselect("Convert values to Title Case in these columns", options=exp_cols, key="map_data_camel")
        with mc2:
            map_padding = st.slider("Extra column-width padding (chars)", 0, 20, 5, key="map_pad")
            map_drop_empty = st.checkbox("Drop fully-empty rows", value=False, key="map_drop")

        st.markdown("---")
        st.markdown("#### Convert column data types")
        st.markdown('<div class="tip-box">Select columns to convert. Phone number keeps only the last 10 digits.</div>', unsafe_allow_html=True)
        map_type_conversions = {}
        map_num_conv = st.number_input("How many columns to convert?", min_value=0, max_value=20, value=0, step=1, key="map_num_conv")
        for i in range(int(map_num_conv)):
            mc1b, mc2b = st.columns([2, 1])
            col_choice = mc1b.selectbox(f"Column #{i+1}", options=["— select —"] + exp_cols, key=f"map_conv_col_{i}")
            type_choice = mc2b.selectbox("Convert to", options=["Integer", "Float", "Phone (last 10 digits)"], key=f"map_conv_type_{i}")
            if col_choice != "— select —":
                map_type_conversions[col_choice] = type_choice

        st.markdown("---")
        st.markdown("#### Sort data")
        map_sort_cols, map_sort_orders = render_sort_ui(exp_cols, key_prefix="map")

        st.markdown("---")
        st.markdown("### 🔄 Preview mapped result")

        out_data = {}
        for exp_col, base_col in mapping.items():
            if base_col and base_col in base_df.columns:
                out_data[exp_col] = base_df[base_col].values
            else:
                out_data[exp_col] = [None] * len(base_df)
        out_df = pd.DataFrame(out_data)

        out_df = apply_transformations(
            out_df, map_drop_empty, map_type_conversions, map_data_camel_cols,
            map_all_camel, map_sort_cols, map_sort_orders
        )

        st.markdown(f"Showing all **{len(out_df)}** rows")
        st.markdown(df_to_html_preview(out_df), unsafe_allow_html=True)

        with st.expander("📋 Mapping summary"):
            summary_rows = [{"Expected Column": ec, "Mapped From (Base)": bc if bc else "— empty —",
                             "Status": "✅ Mapped" if bc else "⬜ Empty"} for ec, bc in mapping.items()]
            st.dataframe(pd.DataFrame(summary_rows), use_container_width=True, hide_index=True)

        st.markdown("---")
        base_name = base_file.name.rsplit('.', 1)[0] + "_mapped"
        download_section(out_df, map_padding, base_name, key_prefix="map")

    elif exp_file or base_file:
        st.info("Please upload **both** the expected schema file and the base data file to continue.")